import os
import pymupdf  # PyMuPDF
from pptx import Presentation
from docx import Document

def extract_text(file_path: str) -> tuple[str, dict]:
    """
    Extracts text from PPTX, PDF, or DOCX files.

    Args:
        file_path: The path to the file to process.

    Returns:
        A tuple of (extracted_text, metadata_dict).
    """
    ext = os.path.splitext(file_path)[1].lower()
    text = ""
    metadata = {"file_path": file_path}

    try:
        if ext == '.pdf':
            doc = pymupdf.open(file_path)
            text_parts = []
            for i, page in enumerate(doc):
                text_parts.append(f"--- Page {i + 1} ---\n{page.get_text()}")
            text = "\n\n".join(text_parts)
            metadata["page_count"] = doc.page_count
            doc.close()

        elif ext == '.pptx':
            prs = Presentation(file_path)
            text_parts = []
            for i, slide in enumerate(prs.slides):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        slide_text.append(shape.text)
                text_parts.append(f"--- Slide {i + 1} ---\n{' '.join(slide_text)}")
            text = "\n\n".join(text_parts)
            metadata["slide_count"] = len(prs.slides)

        elif ext == '.docx':
            doc = Document(file_path)
            text = "\n".join([p.text for p in doc.paragraphs])
            metadata["paragraph_count"] = len(doc.paragraphs)

        else:
            text = f"Unsupported file format: {ext}"
            metadata["error"] = "Unsupported"

    except Exception as e:
        text = f"Error processing file: {str(e)}"
        metadata["error"] = str(e)

    return text, metadata
